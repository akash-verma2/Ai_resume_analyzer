import io
import re
from docx import Document

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract clean text content from DOCX binary stream including tables."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        
        full_text = "\n\n".join(paragraphs)
        full_text = re.sub(r'[ \t]+', ' ', full_text)
        full_text = re.sub(r'\n{3,}', '\n\n', full_text)
        return full_text.strip()
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")
